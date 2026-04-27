#!/usr/bin/env python3
"""
Generate a Terra-compatible sample table TSV (and optional Excel workbook)
for the AFI Rickettsiales pipeline.

Usage
-----
  # Create a blank template:
  python3 make_terra_import_sheet.py --template --output my_run.tsv

  # Validate & export an existing CSV/TSV:
  python3 make_terra_import_sheet.py --input samples.csv --output terra_import.tsv

  # Also write an annotated Excel workbook:
  python3 make_terra_import_sheet.py --input samples.csv --output terra_import.tsv --excel

Terra import format
-------------------
Upload the output TSV in your Terra workspace via:
  Data → Import Data → Upload TSV
The first column header must be "entity:sample_id".

Required columns (in sample table)
-----------------------------------
  sample_id    — unique identifier (alphanumeric, hyphens, underscores)
  run_id       — sequencing run identifier; samples from different runs may
                 appear in the same sheet (multi-run batch support)
  r1_fastq     — GCS path to R1 FASTQ (gs://...)
  r2_fastq     — GCS path to R2 FASTQ (gs://...)
  sample_type  — one of: NTC, NC, PC_MIX8, PC_SINGLE, MIXED4, clinical, PC
  mode         — routine  OR  validation
  expected_taxa — semicolon-delimited expected genera for validation samples;
                  leave blank ("") for routine samples

Validation rules
----------------
  1. Each run_id must contain at least one NTC or NC sample.
  2. Each run_id must contain at least one positive control
     (PC_MIX8, PC_SINGLE, MIXED4, or PC).
  3. Validation-mode samples must have a non-empty expected_taxa.
  4. sample_type and mode values must be from the allowed sets.
"""
import argparse
import csv
import sys
from pathlib import Path

# ---------------------------------------------------------------------------
# Constants
# ---------------------------------------------------------------------------
REQUIRED_COLUMNS = [
    "sample_id", "run_id", "r1_fastq", "r2_fastq",
    "sample_type", "mode", "expected_taxa",
]

VALID_SAMPLE_TYPES = {"NTC", "NC", "PC_MIX8", "PC_SINGLE", "MIXED4", "clinical", "PC"}
VALID_MODES        = {"routine", "validation"}
NTC_TYPES          = {"NTC", "NC"}
PC_TYPES           = {"PC_MIX8", "PC_SINGLE", "MIXED4", "PC"}

TERRA_ENTITY_COLUMN = "entity:sample_id"

# ---------------------------------------------------------------------------
# Template rows (two example runs, each with NTC + PC + clinical)
# ---------------------------------------------------------------------------
TEMPLATE_ROWS = [
    {
        "sample_id":    "run1_NTC_001",
        "run_id":       "run1",
        "r1_fastq":     "gs://your-bucket/run1/NTC_001_R1.fastq.gz",
        "r2_fastq":     "gs://your-bucket/run1/NTC_001_R2.fastq.gz",
        "sample_type":  "NTC",
        "mode":         "routine",
        "expected_taxa": "",
    },
    {
        "sample_id":    "run1_PC_MIX8_001",
        "run_id":       "run1",
        "r1_fastq":     "gs://your-bucket/run1/PC_MIX8_001_R1.fastq.gz",
        "r2_fastq":     "gs://your-bucket/run1/PC_MIX8_001_R2.fastq.gz",
        "sample_type":  "PC_MIX8",
        "mode":         "validation",
        "expected_taxa": "Rickettsia;Orientia;Anaplasma;Ehrlichia;Bartonella;Coxiella;Borrelia;Neoehrlichia",
    },
    {
        "sample_id":    "run1_sample_001",
        "run_id":       "run1",
        "r1_fastq":     "gs://your-bucket/run1/sample_001_R1.fastq.gz",
        "r2_fastq":     "gs://your-bucket/run1/sample_001_R2.fastq.gz",
        "sample_type":  "clinical",
        "mode":         "routine",
        "expected_taxa": "",
    },
    {
        "sample_id":    "run2_NTC_001",
        "run_id":       "run2",
        "r1_fastq":     "gs://your-bucket/run2/NTC_001_R1.fastq.gz",
        "r2_fastq":     "gs://your-bucket/run2/NTC_001_R2.fastq.gz",
        "sample_type":  "NTC",
        "mode":         "routine",
        "expected_taxa": "",
    },
    {
        "sample_id":    "run2_PC_001",
        "run_id":       "run2",
        "r1_fastq":     "gs://your-bucket/run2/PC_001_R1.fastq.gz",
        "r2_fastq":     "gs://your-bucket/run2/PC_001_R2.fastq.gz",
        "sample_type":  "PC",
        "mode":         "routine",
        "expected_taxa": "",
    },
    {
        "sample_id":    "run2_sample_001",
        "run_id":       "run2",
        "r1_fastq":     "gs://your-bucket/run2/sample_001_R1.fastq.gz",
        "r2_fastq":     "gs://your-bucket/run2/sample_001_R2.fastq.gz",
        "sample_type":  "clinical",
        "mode":         "routine",
        "expected_taxa": "",
    },
]


# ---------------------------------------------------------------------------
# I/O helpers
# ---------------------------------------------------------------------------

def read_input(path: str) -> list[dict]:
    """Read CSV or TSV (auto-detected by extension)."""
    p = Path(path)
    delimiter = "\t" if p.suffix.lower() in (".tsv", ".txt") else ","
    with open(p, encoding="utf-8-sig") as fh:
        reader = csv.DictReader(fh, delimiter=delimiter)
        return [dict(row) for row in reader]


def write_tsv(rows: list[dict], path: str) -> None:
    """Write Terra-compatible TSV with entity:sample_id as first column."""
    fieldnames = [TERRA_ENTITY_COLUMN] + [c for c in REQUIRED_COLUMNS if c != "sample_id"]
    with open(path, "w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=fieldnames, delimiter="\t",
                                extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            terra_row = {k: v for k, v in row.items()}
            terra_row[TERRA_ENTITY_COLUMN] = row.get("sample_id", "")
            writer.writerow(terra_row)


# ---------------------------------------------------------------------------
# Validation
# ---------------------------------------------------------------------------

def validate(rows: list[dict]) -> list[str]:
    """Return list of error messages; empty list = all good."""
    errors: list[str] = []

    # Check required columns
    if not rows:
        return ["Input file is empty."]
    missing_cols = [c for c in REQUIRED_COLUMNS if c not in rows[0]]
    if missing_cols:
        errors.append(f"Missing required columns: {', '.join(missing_cols)}")
        return errors  # can't proceed without columns

    # Per-row validation
    for i, row in enumerate(rows, start=2):  # row 2 = first data row
        sid  = row.get("sample_id", "").strip()
        st   = row.get("sample_type", "").strip()
        mode = row.get("mode", "").strip()
        exp  = row.get("expected_taxa", "").strip()

        if not sid:
            errors.append(f"Row {i}: sample_id is empty.")
        if st not in VALID_SAMPLE_TYPES:
            errors.append(
                f"Row {i} ({sid}): invalid sample_type '{st}'. "
                f"Allowed: {', '.join(sorted(VALID_SAMPLE_TYPES))}"
            )
        if mode not in VALID_MODES:
            errors.append(
                f"Row {i} ({sid}): invalid mode '{mode}'. "
                f"Allowed: {', '.join(sorted(VALID_MODES))}"
            )
        if mode == "validation" and not exp:
            errors.append(
                f"Row {i} ({sid}): mode=validation but expected_taxa is empty."
            )

    # Per-run validation
    from collections import defaultdict
    run_types: dict[str, set] = defaultdict(set)
    for row in rows:
        rid = row.get("run_id", "").strip()
        st  = row.get("sample_type", "").strip().upper()
        run_types[rid].add(st)

    for rid, types in run_types.items():
        if not types & {t.upper() for t in NTC_TYPES}:
            errors.append(
                f"run_id '{rid}': no NTC or NC sample found. "
                "Every run must have a negative template control."
            )
        if not types & {t.upper() for t in PC_TYPES}:
            errors.append(
                f"run_id '{rid}': no positive control found "
                f"(need one of: {', '.join(sorted(PC_TYPES))})."
            )

    return errors


# ---------------------------------------------------------------------------
# Excel export (optional — requires openpyxl)
# ---------------------------------------------------------------------------

def write_excel(rows: list[dict], path: str) -> None:
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError:
        print(
            "WARNING: openpyxl not installed; skipping Excel export.\n"
            "Install it with:  pip install openpyxl",
            file=sys.stderr,
        )
        return

    wb = openpyxl.Workbook()

    # ── Sheet 1: Data ────────────────────────────────────────────────────────
    ws = wb.active
    ws.title = "samples"

    header_fill   = PatternFill("solid", fgColor="1F4E79")
    header_font   = Font(bold=True, color="FFFFFF")
    req_fill      = PatternFill("solid", fgColor="D6E4F0")
    ntc_fill      = PatternFill("solid", fgColor="E2EFDA")   # green tint
    pc_fill       = PatternFill("solid", fgColor="FFF2CC")   # yellow tint

    fieldnames = [TERRA_ENTITY_COLUMN] + [c for c in REQUIRED_COLUMNS if c != "sample_id"]
    for col_idx, name in enumerate(fieldnames, start=1):
        cell = ws.cell(row=1, column=col_idx, value=name)
        cell.font  = header_font
        cell.fill  = header_fill
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    for row_idx, row in enumerate(rows, start=2):
        st = row.get("sample_type", "").strip().upper()
        row_fill = ntc_fill if st in {t.upper() for t in NTC_TYPES} else (
                   pc_fill  if st in {t.upper() for t in PC_TYPES}  else None)

        for col_idx, col_name in enumerate(fieldnames, start=1):
            key = "sample_id" if col_name == TERRA_ENTITY_COLUMN else col_name
            cell = ws.cell(row=row_idx, column=col_idx, value=row.get(key, ""))
            cell.fill = req_fill
            if row_fill:
                cell.fill = row_fill

    # Auto-size columns
    for col_idx, col_name in enumerate(fieldnames, start=1):
        max_len = max(
            len(str(ws.cell(row=r, column=col_idx).value or ""))
            for r in range(1, ws.max_row + 1)
        )
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 60)

    ws.freeze_panes = "A2"

    # ── Sheet 2: Instructions ────────────────────────────────────────────────
    inst = wb.create_sheet("instructions")
    inst.column_dimensions["A"].width = 20
    inst.column_dimensions["B"].width = 70

    inst_rows = [
        ("Column",          "Description"),
        ("entity:sample_id","Unique sample identifier (alphanumeric, hyphens, underscores)."),
        ("run_id",          "Sequencing run ID. Include multiple runs in one sheet to batch them together."),
        ("r1_fastq",        "GCS path to R1 FASTQ file (gs://bucket/path/file_R1.fastq.gz)."),
        ("r2_fastq",        "GCS path to R2 FASTQ file (gs://bucket/path/file_R2.fastq.gz)."),
        ("sample_type",     f"One of: {', '.join(sorted(VALID_SAMPLE_TYPES))}"),
        ("mode",            "routine  — no expected-taxa check.\nvalidation  — compares detected vs expected_taxa."),
        ("expected_taxa",   "Required for validation mode. Semicolon-separated genera, e.g. Rickettsia;Orientia"),
        ("",                ""),
        ("RULES",           ""),
        ("NTC/NC required", "Each run_id must have at least one NTC or NC sample."),
        ("PC required",     f"Each run_id must have at least one positive control: {', '.join(sorted(PC_TYPES))}"),
        ("",                ""),
        ("COLOR KEY",       ""),
        ("Green rows",      "NTC / NC samples"),
        ("Yellow rows",     "Positive control samples (PC_MIX8, PC_SINGLE, MIXED4, PC)"),
        ("Blue rows",       "Clinical / routine samples"),
        ("",                ""),
        ("IMPORT",          "In Terra: Data → Import Data → Upload TSV → select terra_import.tsv"),
    ]

    title_font  = Font(bold=True, color="1F4E79")
    normal_font = Font()
    for r_idx, (label, desc) in enumerate(inst_rows, start=1):
        c1 = inst.cell(row=r_idx, column=1, value=label)
        c2 = inst.cell(row=r_idx, column=2, value=desc)
        if label in ("Column", "RULES", "COLOR KEY", "IMPORT"):
            c1.font = title_font
            c2.font = title_font
        c2.alignment = Alignment(wrap_text=True)

    inst.row_dimensions[1].height = 20

    wb.save(path)
    print(f"Excel workbook written: {path}", file=sys.stderr)


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main() -> None:
    parser = argparse.ArgumentParser(
        description="Generate or validate a Terra sample import sheet for the AFI Rickettsiales pipeline.",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=__doc__,
    )
    parser.add_argument("--input",    "-i", help="Existing CSV or TSV to validate and convert.")
    parser.add_argument("--output",   "-o", default="terra_import.tsv",
                        help="Output TSV path (default: terra_import.tsv).")
    parser.add_argument("--excel",    "-x", action="store_true",
                        help="Also write an annotated Excel workbook (.xlsx).")
    parser.add_argument("--template", "-t", action="store_true",
                        help="Write a blank template with example rows instead of validating input.")
    args = parser.parse_args()

    if args.template:
        rows = TEMPLATE_ROWS
        print(f"Writing template with {len(rows)} example rows …", file=sys.stderr)
    elif args.input:
        rows = read_input(args.input)
        print(f"Read {len(rows)} rows from {args.input}", file=sys.stderr)
    else:
        parser.error("Provide --input <file> or --template.")
        return

    # Validate
    errors = validate(rows)
    if errors:
        print("\nValidation errors:", file=sys.stderr)
        for err in errors:
            print(f"  ✗ {err}", file=sys.stderr)
        if not args.template:
            sys.exit(1)
    else:
        print("Validation passed.", file=sys.stderr)

    # Write TSV
    write_tsv(rows, args.output)
    print(f"Terra import TSV written: {args.output}", file=sys.stderr)

    # Write Excel
    if args.excel:
        xl_path = str(Path(args.output).with_suffix(".xlsx"))
        write_excel(rows, xl_path)


if __name__ == "__main__":
    main()
